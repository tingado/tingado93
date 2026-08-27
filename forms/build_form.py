import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

wb = openpyxl.Workbook()

# ---------- survey sheet ----------
survey = wb.active
survey.title = "survey"

survey_headers = [
    "type", "name", "label", "hint", "required", "default",
    "appearance", "relevant", "constraint", "constraint_message",
]
survey.append(survey_headers)

rows = [
    ("select_one tipo_registro", "tipo_registro", "Tipo de Registro", "Seleccione si es el registro de apertura o de cierre de bodega", "yes", "", "", "", "", ""),
    ("text", "encargado", "Encargado de la información", "Nombre completo de quien diligencia el registro", "yes", "", "", "", "", ""),
    ("date", "fecha_info", "Fecha de la información", "", "yes", "today()", "", "", "", ""),
    ("time", "hora_info", "Hora", "", "yes", "", "", "", "", ""),
    ("geopoint", "ubicacion", "Ubicación de la bodega (GPS)", "Capture la ubicación GPS actual; se exporta como KML/KMZ desde KoboToolbox (Descargar > GPS/KML)", "yes", "", "", "", "", ""),
    ("text", "comentario_general", "Comentario General", "Observaciones adicionales sobre el inventario", "no", "", "multiline", "", "", ""),
    ("begin group", "fotos", "Fotos de Referencia del Inventario", "", "", "", "field-list", "", "", ""),
    ("image", "foto_1", "Foto de referencia 1", "", "yes", "", "", "", "", ""),
    ("image", "foto_2", "Foto de referencia 2", "", "yes", "", "", "", "", ""),
    ("image", "foto_3", "Foto de referencia 3", "", "yes", "", "", "", "", ""),
    ("image", "foto_4", "Foto de referencia 4", "", "yes", "", "", "", "", ""),
    ("image", "foto_5", "Foto de referencia 5", "", "yes", "", "", "", "", ""),
    ("image", "foto_6", "Foto de referencia 6", "", "yes", "", "", "", "", ""),
    ("end group", "", "", "", "", "", "", "", "", ""),
]
for r in rows:
    survey.append(r)

# ---------- choices sheet ----------
choices = wb.create_sheet("choices")
choices.append(["list_name", "name", "label"])
choices.append(["tipo_registro", "inicial", "Registro inicial de Bodega"])
choices.append(["tipo_registro", "cierre", "Registro cierre día Bodega"])

# ---------- settings sheet ----------
settings = wb.create_sheet("settings")
settings.append(["form_title", "form_id", "version", "default_language"])
settings.append(["Inventario Fotográfico de Bodega", "inventario_fotografico_bodega", "1", "Español (es)"])

# ---------- formatting ----------
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
body_font = Font(name="Arial")

for sheet in (survey, choices, settings):
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
            if cell.row == 1:
                cell.fill = header_fill
    for col_cells in sheet.columns:
        length = max((len(str(c.value)) if c.value else 0) for c in col_cells)
        sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 12), 45)
    sheet.freeze_panes = "A2"

wb.save("forms/inventario_fotografico_bodega.xlsx")
print("saved")
